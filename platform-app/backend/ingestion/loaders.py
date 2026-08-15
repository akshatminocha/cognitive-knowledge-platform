"""
File Loaders — Parse various file formats into plain text.

Supports: PDF, TXT, Markdown, CSV, JSON, DOCX
All loaders return a LoadedDocument with the extracted text content.
"""

from __future__ import annotations

import csv
import io
import json
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


@dataclass
class LoadedDocument:
    """Result of loading a file."""

    content: str
    filename: str
    file_type: str
    page_count: int = 1
    metadata: dict = field(default_factory=dict)


class FileLoader:
    """
    Multi-format file loader.

    Dispatches to the appropriate parser based on file extension.

    Usage:
        loader = FileLoader()
        doc = loader.load(Path("report.pdf"))
        print(doc.content[:200])
    """

    SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".csv", ".json", ".docx"}

    def load(self, file_path: Path) -> LoadedDocument:
        """Load a file and return its text content."""
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = file_path.suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type: {ext}. "
                f"Supported: {self.SUPPORTED_EXTENSIONS}"
            )

        loader_map = {
            ".txt": self._load_text,
            ".md": self._load_text,
            ".csv": self._load_csv,
            ".json": self._load_json,
            ".pdf": self._load_pdf,
            ".docx": self._load_docx,
        }

        loader_fn = loader_map[ext]
        content = loader_fn(file_path)

        logger.info(f"Loaded {file_path.name}: {len(content)} chars")

        return LoadedDocument(
            content=content,
            filename=file_path.name,
            file_type=ext.lstrip("."),
        )

    def load_bytes(self, content: bytes, filename: str) -> LoadedDocument:
        """Load from raw bytes (for API file uploads)."""
        ext = Path(filename).suffix.lower()

        if ext in (".txt", ".md"):
            text = content.decode("utf-8", errors="replace")
        elif ext == ".csv":
            text = self._parse_csv_bytes(content)
        elif ext == ".json":
            text = self._parse_json_bytes(content)
        elif ext == ".pdf":
            text = self._parse_pdf_bytes(content)
        elif ext == ".docx":
            text = self._parse_docx_bytes(content)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

        return LoadedDocument(
            content=text,
            filename=filename,
            file_type=ext.lstrip("."),
        )

    # --- Text/Markdown ---
    def _load_text(self, path: Path) -> str:
        return path.read_text(encoding="utf-8", errors="replace")

    # --- CSV ---
    def _load_csv(self, path: Path) -> str:
        return self._parse_csv_bytes(path.read_bytes())

    def _parse_csv_bytes(self, content: bytes) -> str:
        """Convert CSV to a readable text representation."""
        text = content.decode("utf-8", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        rows = []
        for row in reader:
            row_text = " | ".join(f"{k}: {v}" for k, v in row.items() if v)
            rows.append(row_text)
        return "\n".join(rows)

    # --- JSON ---
    def _load_json(self, path: Path) -> str:
        return self._parse_json_bytes(path.read_bytes())

    def _parse_json_bytes(self, content: bytes) -> str:
        """Convert JSON to readable text."""
        data = json.loads(content.decode("utf-8", errors="replace"))

        if isinstance(data, list):
            # Array of objects → one row per object
            rows = []
            for item in data:
                if isinstance(item, dict):
                    row_text = " | ".join(f"{k}: {v}" for k, v in item.items())
                    rows.append(row_text)
                else:
                    rows.append(str(item))
            return "\n".join(rows)
        elif isinstance(data, dict):
            return json.dumps(data, indent=2)
        else:
            return str(data)

    # --- PDF ---
    def _load_pdf(self, path: Path) -> str:
        return self._parse_pdf_bytes(path.read_bytes())

    def _parse_pdf_bytes(self, content: bytes) -> str:
        """Extract text from PDF using PyMuPDF (fitz)."""
        try:
            import fitz  # PyMuPDF

            doc = fitz.open(stream=content, filetype="pdf")
            pages = []
            for page in doc:
                pages.append(page.get_text())
            doc.close()
            return "\n\n".join(pages)
        except ImportError:
            logger.warning(
                "PyMuPDF (fitz) not installed. Install with: pip install pymupdf"
            )
            return "[PDF parsing requires PyMuPDF: pip install pymupdf]"

    # --- DOCX ---
    def _load_docx(self, path: Path) -> str:
        return self._parse_docx_bytes(path.read_bytes())

    def _parse_docx_bytes(self, content: bytes) -> str:
        """Extract text from DOCX using python-docx."""
        try:
            import docx

            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n\n".join(paragraphs)
        except ImportError:
            logger.warning(
                "python-docx not installed. Install with: pip install python-docx"
            )
            return "[DOCX parsing requires python-docx: pip install python-docx]"
