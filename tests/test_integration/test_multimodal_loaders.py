"""
Tests for Multimodal Ingestion Loaders (.pptx, .png, .jpg, .wav, .mp3, .mp4).
"""

import io
import pytest
from pathlib import Path
from PIL import Image
import pptx

from backend.ingestion.loaders import FileLoader, LoadedDocument
from backend.ingestion.pipeline import IngestionPipeline


class TestMultimodalLoaders:
    """Test suite for multimodal file loaders."""

    def test_pptx_loader(self):
        """Test PPTX presentation loading with slides, text frames, and speaker notes."""
        prs = pptx.Presentation()
        # Add slide 1 with title
        slide1 = prs.slides.add_slide(prs.slide_layouts[0])
        slide1.shapes.title.text = "Cardiovascular Health Overview"
        # Add speaker note
        slide1.notes_slide.notes_text_frame.text = "Dr. Smith presented findings on patient cohort."

        # Add slide 2 with table
        slide2 = prs.slides.add_slide(prs.slide_layouts[5])
        rows, cols = 2, 2
        left, top, width, height = 100, 100, 400, 200
        table_shape = slide2.shapes.add_table(rows, cols, left, top, width, height)
        table = table_shape.table
        table.cell(0, 0).text = "Patient ID"
        table.cell(0, 1).text = "Status"
        table.cell(1, 0).text = "PAT-101"
        table.cell(1, 1).text = "Stable"

        buf = io.BytesIO()
        prs.save(buf)
        pptx_bytes = buf.getvalue()

        loader = FileLoader()
        doc = loader.load_bytes(pptx_bytes, "clinical_presentation.pptx")

        assert isinstance(doc, LoadedDocument)
        assert doc.file_type == "pptx"
        assert doc.page_count == 2
        assert "Cardiovascular Health Overview" in doc.content
        assert "Dr. Smith presented findings" in doc.content
        assert "PAT-101" in doc.content
        assert doc.metadata.get("multimodal_type") == "presentation"

    def test_vision_image_loader_png(self):
        """Test PNG vision document parsing with dimensions and metadata."""
        img = Image.new("RGB", (320, 240), color="teal")
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        png_bytes = buf.getvalue()

        loader = FileLoader()
        doc = loader.load_bytes(png_bytes, "cardiology_scan.png")

        assert doc.file_type == "png"
        assert doc.page_count == 1
        assert doc.metadata.get("multimodal_type") == "vision"
        assert len(doc.content.strip()) > 0

    def test_vision_image_loader_jpeg(self):
        """Test JPEG vision document parsing."""
        img = Image.new("RGB", (640, 480), color="purple")
        buf = io.BytesIO()
        img.save(buf, format="JPEG")
        jpg_bytes = buf.getvalue()

        loader = FileLoader()
        doc = loader.load_bytes(jpg_bytes, "clinical_chart.jpg")

        assert doc.file_type == "jpg"
        assert doc.metadata.get("multimodal_type") == "vision"
        assert len(doc.content.strip()) > 0

    def test_audio_loader_wav_and_mp3(self):
        """Test audio recording parsing with speaker diarization structure."""
        loader = FileLoader()
        fake_wav = b"RIFF" + b"\x00" * 200 + b"WAVEfmt "
        doc_wav = loader.load_bytes(fake_wav, "patient_interview.wav")
        assert doc_wav.file_type == "wav"
        assert doc_wav.metadata.get("multimodal_type") == "audio"
        assert "Speaker" in doc_wav.content or "Audio" in doc_wav.content

        fake_mp3 = b"\xFF\xFB" + b"\x00" * 300
        doc_mp3 = loader.load_bytes(fake_mp3, "doctor_dictation.mp3")
        assert doc_mp3.file_type == "mp3"
        assert doc_mp3.metadata.get("multimodal_type") == "audio"
        assert "Speaker" in doc_mp3.content or "Audio" in doc_mp3.content

    def test_video_loader_mp4(self):
        """Test video chronological scene summarization."""
        loader = FileLoader()
        fake_mp4 = b"\x00\x00\x00\x20ftypisom" + b"\x00" * 500
        doc_mp4 = loader.load_bytes(fake_mp4, "surgical_demonstration.mp4")
        assert doc_mp4.file_type == "mp4"
        assert doc_mp4.metadata.get("multimodal_type") == "video"
        assert "Video Understanding" in doc_mp4.content
        assert "[Scene 1" in doc_mp4.content


class TestMultimodalIngestionPipeline:
    """Integration test: pipeline ingestion of multimodal files."""

    @pytest.mark.asyncio
    async def test_pipeline_ingests_pptx_deck(self):
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Oncology Clinical Trial Protocol"
        slide.notes_slide.notes_text_frame.text = "Phase 2 multi-center investigation into immunotherapy."

        buf = io.BytesIO()
        prs.save(buf)
        pptx_bytes = buf.getvalue()

        pipeline = IngestionPipeline(schema_name="healthtech")
        res = await pipeline.ingest_bytes(pptx_bytes, "trial_protocol.pptx")

        assert res.filename == "trial_protocol.pptx"
        assert res.status == "success"
        assert res.chunks_created > 0
